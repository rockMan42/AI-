import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HEADERS = (
    "维度",
    "平均分",
    "标准差",
    "最低分",
    "最高分",
    "P25",
    "P50",
    "P75",
    "人数",
)


def build_performance_workbook(
    cycle: str,
    dimension: str,
    rows: list[dict],
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "绩效统计"
    sheet.append([f"{cycle} 绩效统计", f"维度：{dimension}"])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.append(HEADERS)

    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        sheet.append([
            row["dimension_value"],
            row["avg_score"],
            row["std_dev"],
            row["min_score"],
            row["max_score"],
            row["p25"],
            row["p50"],
            row["p75"],
            row["total_count"],
        ])

    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:I{max(sheet.max_row, 2)}"
    widths = (24, 12, 12, 12, 12, 10, 10, 10, 10)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
